"""E2E (modo demo, Chromium) · Informe Cristian con criterio de HOJAS PUBLICADAS: pantalla = Excel descargado = informe generado.
Datos ficticios. Una hoja de Cristian publicada HOY con todas sus entregas pendientes debe aparecer.
Código de salida: 0 todo PASS · 1 alguna FAIL · 2 error de preparación o casos faltantes."""
import asyncio, os, subprocess, sys, time, socket, traceback, re
from playwright.async_api import async_playwright
from openpyxl import load_workbook
ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..'))
DIST = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, 'dist-demo'); OUT = '/tmp/okd_informe_e2e'; PORT = 4381; R = []; EXPECTED = 5
def res(n, ok, d=''): R.append((n, 'PASS' if ok else 'FAIL', str(d)[:260])); print('..', ' | '.join(R[-1]), flush=True)
def num(t): return int(re.sub(r'[^0-9]', '', t) or 0)
PUB = '''async()=>{const A=window.OKD_ONLINE.api; const mk=(n,z,fee)=>({ready:true, client_name:'CLIENTE INFORME '+n, client_key:'qa-inf-'+n, address:'Calle '+n, zone:z, order_numbers:[String(n)], merchandise_amount:250000, shipping_fee:fee, bultos:null, task:'ENTREGA', stop_type:'PEDIDO'});
  const r = await A.publishRoute({sheet_key:'qa-inf-'+Date.now(), rules_version:1, routing_rules_version:1, carrier_id:'CT-ESTE', responsable_id:'u-cristian', stops:[mk(85001,'ESTE',13000), mk(85002,'ESTE',13000), mk(85003,'CENTRO',7000)]});
  return {id:r.id, number:r.number, estados:r.stops.map(s=>s.outcome)};}'''
async def main():
    if not os.path.isfile(os.path.join(DIST, 'index.html')): raise RuntimeError('falta el build demo en ' + DIST)
    os.makedirs(OUT, exist_ok=True)
    srv = subprocess.Popen([sys.executable, '-m', 'http.server', str(PORT), '--bind', '127.0.0.1'], cwd=DIST, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    try:
        for _ in range(60):
            try: socket.create_connection(('127.0.0.1', PORT), 1).close(); break
            except OSError: time.sleep(0.2)
        async with async_playwright() as p:
          b = await p.chromium.launch(executable_path=os.environ.get('CHROMIUM_PATH') or None)
          try:
            ctx = await b.new_context(viewport={'width': 1366, 'height': 900}, accept_downloads=True); pg = await ctx.new_page(); pg.on('dialog', lambda d: asyncio.ensure_future(d.accept()))
            await pg.goto(f'http://127.0.0.1:{PORT}/'); await pg.evaluate('()=>localStorage.clear()'); await pg.goto(f'http://127.0.0.1:{PORT}/'); await pg.wait_for_timeout(600)
            await pg.fill('#loginEmail', 'admin@okdespacho.demo'); await pg.fill('#loginPass', 'demo'); await pg.click('#loginBtn'); await pg.wait_for_timeout(2500)
            if await pg.locator('#pinA').count() and await pg.is_visible('#pinA'): await pg.fill('#pinA', '123456'); await pg.fill('#pinB', '123456'); await pg.click('#pinGo'); await pg.wait_for_timeout(900)
            # informe ANTES de publicar la hoja nueva (7 días)
            await pg.evaluate('()=>{location.hash="#/informes"}'); await pg.wait_for_timeout(2500)
            t0 = await pg.locator('#rpTotal').inner_text()
            hoja = await pg.evaluate(PUB)
            await pg.evaluate('()=>{location.hash="#/"}'); await pg.wait_for_timeout(800); await pg.evaluate('()=>{location.hash="#/informes"}'); await pg.wait_for_timeout(2500)
            t1 = await pg.locator('#rpTotal').inner_text(); crit = await pg.locator('#rpCriterion').inner_text(); states = await pg.locator('#rpStates').inner_text()
            n0 = num(re.search(r'([\d.]+)\s*envíos', t0).group(1)); n1 = num(re.search(r'([\d.]+)\s*envíos', t1).group(1))
            m0 = num(re.findall(r'\$\s*[\d.]+', t0)[0]); m1 = num(re.findall(r'\$\s*[\d.]+', t1)[0])
            res('INF-01 hoja de Cristian publicada con TODAS las entregas pendientes aparece: +3 envíos y +$33.000 (sin que Cristian la complete)',
                all(e == 'PENDIENTE' for e in hoja['estados']) and n1 - n0 == 3 and m1 - m0 == 33000, {'antes': t0, 'despues': t1, 'hoja': hoja['number']})
            res('INF-02 la pantalla explica el criterio (hojas publicadas por fecha de publicación) y no presenta pendientes como entregados',
                'PUBLICADAS' in crit and 'publicación' in crit and 'pendientes' in states and 'no es dinero cobrado' in states and 'Hojas finalizadas' not in await pg.locator('body').inner_text(), {'criterio': crit[:120], 'estados': states[:160]})
            # el aviso de notificaciones se cierra con "Más tarde", como lo haría la persona
            for _ in range(3):
                if await pg.locator('#pushLater').count() and await pg.is_visible('#pushLater'): await pg.click('#pushLater'); await pg.wait_for_timeout(300)
            # Excel descargado = pantalla
            async with pg.expect_download() as dl: await pg.click('#rpExcel')
            f = os.path.join(OUT, 'informe_cristian.xlsx'); await (await dl.value).save_as(f)
            ws = load_workbook(f)['Resumen']; rows = [[c.value for c in r] for r in ws.iter_rows()]
            tot = next(r for r in rows if r and r[0] == 'TOTAL'); head = next(r for r in rows if r and r[0] == 'ZONA')
            res('INF-03 el Excel descargado coincide con la pantalla (envíos e importe) y usa los rótulos nuevos',
                tot[1] == n1 and round(tot[6]) == m1 and head[1] == 'ENVÍOS INCLUIDOS' and 'IMPORTE DE ENVÍO INFORMADO' in head, {'excel_total': tot, 'pantalla': [n1, m1]})
            # Generar → queda en el historial como V2 con los mismos números
            await pg.click('#rpGenerate'); await pg.wait_for_timeout(3000)
            # mismas llamadas que la pantalla de historial: history → get(report_id)
            hist = await pg.evaluate('''async()=>{ const l=(await window.OKD_ONLINE.api.reports('history', { filters: {} })).reports || []; const last=l.sort((a,b)=>String(b.generated_at).localeCompare(String(a.generated_at)))[0];
              if(!last) return null; const s=(await window.OKD_ONLINE.api.reports('get', { report_id: last.id })).report.snapshot; return {v:s.version, crit:s.criterion, n:s.total.shipments, m:s.total.income} }''')
            res('INF-04 el informe GENERADO queda guardado con el criterio de hojas publicadas y los mismos números de la pantalla', hist and hist['crit'] == 'HOJAS_PUBLICADAS' and hist['n'] == n1 and round(hist['m']) == m1, hist)
            # marcar luego una entrega como completada: totales iguales (sin duplicar ni mover de período)
            await pg.evaluate(f'''()=>{{const k="okd_online_demo_v2",d=JSON.parse(localStorage.getItem(k)); const r=d.routes2.find(x=>x.id==="{hoja['id']}"); r.stops[0].outcome="ENTREGADO"; r.stops[1].outcome="ENTREGA_PARCIAL"; r.operational_status="COMPLETADA"; r.completed_at=new Date(Date.now()+86400000*3).toISOString(); localStorage.setItem(k,JSON.stringify(d));}}''')
            await pg.evaluate('()=>{location.hash="#/"}'); await pg.wait_for_timeout(800); await pg.evaluate('()=>{location.hash="#/informes"}'); await pg.wait_for_timeout(2500)
            t2 = await pg.locator('#rpTotal').inner_text(); n2 = num(re.search(r'([\d.]+)\s*envíos', t2).group(1)); m2 = num(re.findall(r'\$\s*[\d.]+', t2)[0])
            res('INF-05 al completar luego entregas, los totales NO se duplican ni cambian', n2 == n1 and m2 == m1, {'antes': t1, 'despues': t2})
          finally: await b.close()
    finally:
        if srv.poll() is None: srv.terminate(); srv.wait(5)
    print('TOTAL', sum(1 for r in R if r[1] == 'PASS'), 'PASS /', sum(1 for r in R if r[1] == 'FAIL'), 'FAIL · esperados', EXPECTED)
    return R
if __name__ == '__main__':
    try: rr = asyncio.run(main())
    except Exception as e: traceback.print_exc(); print('ERROR:', repr(e)); sys.exit(2)
    if len(rr) != EXPECTED: print('ERROR: se ejecutaron %d de %d' % (len(rr), EXPECTED)); sys.exit(2)
    sys.exit(1 if any(r[1] != 'PASS' for r in rr) else 0)
