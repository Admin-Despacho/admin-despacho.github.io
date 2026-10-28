"""E2E (modo demo, Chromium) · REENVÍOS: el constructor permite volver a cargar un pedido NO ENTREGADO/REPROGRAMADO (misma función de reserva
que usa al agregar pedidos) y el informe de Cristian cuenta el envío UNA vez (pantalla = Excel). Pedidos FICTICIOS.
Exit 0 todo PASS · 1 alguna FAIL · 2 error de preparación o casos faltantes."""
import asyncio, os, sys, re, traceback
sys.path.insert(0, os.path.dirname(__file__))
import run_cambiar_rol as cr
from playwright.async_api import async_playwright
from openpyxl import load_workbook
cr.PORT = 4385; K = cr.K; R = []; EXPECTED = 7; OUT = '/tmp/okd_reenvios'
def res(n, ok, d=''): R.append((n, 'PASS' if ok else 'FAIL', str(d)[:260])); print('..', ' | '.join(R[-1]), flush=True)
def num(t): return int(re.sub(r'[^0-9]', '', t) or 0)
PUB = '''async(a)=>{ const [key, orders, fee, zone, client] = a; const r = await window.OKD_ONLINE.api.publishRoute({sheet_key:key, rules_version:1, routing_rules_version:1, carrier_id:'CT-ESTE', responsable_id:'u-cristian',
  stops: orders.map((o,i)=>({ready:true, client_name: client[i], client_key: client[i], address:'Calle '+o, zone: zone[i], order_numbers:[o], merchandise_amount:200000, shipping_fee: fee[i], task:'ENTREGA', stop_type:'PEDIDO'}))});
  return {id:r.id, number:r.number}; }'''
async def screen(pg):
    await pg.evaluate('()=>{location.hash="#/"}'); await pg.wait_for_timeout(600); await pg.evaluate('()=>{location.hash="#/informes"}'); await pg.wait_for_timeout(2500); await cr.dismiss(pg)
    t = await pg.locator('#rpTotal').inner_text(); return num(re.search(r'([\d.]+)\s*envíos', t).group(1)), num(re.findall(r'\$\s*[\d.]+', t)[0])
async def main():
    os.makedirs(OUT, exist_ok=True); srv = cr.serve()
    try:
        async with async_playwright() as p:
          b = await p.chromium.launch(executable_path=os.environ.get('CHROMIUM_PATH') or None)
          try:
            ctx = await b.new_context(viewport={'width': 1366, 'height': 900}, accept_downloads=True); pg = await ctx.new_page(); pg.on('dialog', lambda d: asyncio.ensure_future(d.accept()))
            await pg.goto(f'http://127.0.0.1:{cr.PORT}/'); await pg.evaluate('()=>localStorage.clear()'); await cr.login(pg, 'admin@okdespacho.demo')
            n0, m0 = await screen(pg)
            # primer intento (hoja de Cristian): RE-1001 NO ENTREGADO, RE-1002 pendiente, RE-1003 ENTREGADO
            r1 = await pg.evaluate(PUB, ['re-h1', ['RE-1001', 'RE-1002', 'RE-1003'], [13000, 7000, 13000], ['ESTE', 'CENTRO', 'ESTE'], ['Cliente Ausente', 'Cliente Pendiente', 'Cliente OK']])
            await pg.evaluate(f'''()=>{{ const d=JSON.parse(localStorage.getItem("{K}")); const r=d.routes2.find(x=>x.id==="{r1['id']}"); r.operational_status='EN_CURSO';
              r.stops.find(s=>s.order_numbers.includes('RE-1001')).outcome='NO_ENTREGADO'; r.stops.find(s=>s.order_numbers.includes('RE-1003')).outcome='ENTREGADO'; localStorage.setItem("{K}", JSON.stringify(d)); }}''')
            n1, m1 = await screen(pg)
            res('RE-E2E-01 la primera hoja suma sus 3 envíos y su importe (33.000)', n1 - n0 == 3 and m1 - m0 == 33000, {'antes': [n0, m0], 'despues': [n1, m1]})
            # ===== constructor: misma función de reserva que usa al agregar pedidos (OKD_BRIDGE.orders.reserve) =====
            await pg.evaluate('()=>{location.hash="#/despacho"}'); await pg.wait_for_timeout(2500); await cr.dismiss(pg)
            await pg.evaluate('async()=>{ await SheetTabs.create(); }'); await pg.wait_for_timeout(1200)
            tries = await pg.evaluate('''async()=>{ const ob=window.OKD_BRIDGE.orders, k=RouteStore.sheetId, out={};
              for (const o of ['RE-1001','RE-1002','RE-1003']) { try { await ob.reserve(k, [o]); out[o]='permitido'; } catch(e) { out[o]=e.message; } } return {k, out}; }''')
            hist = await pg.evaluate(f'()=>(JSON.parse(localStorage.getItem("{K}")).reservationHistory||[]).map(h=>[h.order_id, h.reason, h.number||null])')
            old = await pg.evaluate(f'()=>JSON.parse(localStorage.getItem("{K}")).routes2.find(r=>r.id==="{r1["id"]}").stops.map(s=>[s.order_numbers[0], s.outcome])')
            res('RE-E2E-02 constructor: el pedido NO ENTREGADO se puede volver a cargar en una hoja nueva; queda el historial y la hoja anterior intacta',
                tries['out']['RE-1001'] == 'permitido' and any(h[0] == 'RE-1001' and h[1] == 'REINTENTO_NO_ENTREGADO_O_REPROGRAMADO' for h in hist)
                and old == [['RE-1001', 'NO_ENTREGADO'], ['RE-1002', 'PENDIENTE'], ['RE-1003', 'ENTREGADO']], {'reservas': tries['out'], 'historial': hist})
            res('RE-E2E-03 el pendiente en otra hoja activa y el ya entregado NO se liberan', 'PEDIDO_YA_EN_OTRA_HOJA' in tries['out']['RE-1002'] and 'PEDIDO_YA_EN_OTRA_HOJA' in tries['out']['RE-1003'], tries['out'])
            # segundo intento publicado AL DÍA SIGUIENTE, en OTRA zona y ENTREGADO
            r2 = await pg.evaluate(PUB, [tries['k'], ['RE-1001'], [7000], ['CENTRO'], ['Cliente Ausente']])
            await pg.evaluate(f'''()=>{{ const d=JSON.parse(localStorage.getItem("{K}")); const r=d.routes2.find(x=>x.id==="{r2['id']}"); r.operational_status='COMPLETADA';
              r.published_at=new Date(Date.parse(d.routes2.find(x=>x.id==="{r1['id']}").published_at)+86400000).toISOString(); r.stops[0].outcome='ENTREGADO'; localStorage.setItem("{K}", JSON.stringify(d)); }}''')
            n2, m2 = await screen(pg); retr = await pg.locator('#rpRetries').inner_text() if await pg.locator('#rpRetries').count() else ''
            res('RE-E2E-04 informe: el reintento NO suma otro envío ni otro importe (ni con otra zona e importe)', (n2, m2) == (n1, m1), {'antes': [n1, m1], 'despues': [n2, m2]})
            res('RE-E2E-05 la pantalla muestra el reenvío con sus dos intentos (hoja y resultado) y lo cuenta una vez', 'RE-1001' in retr and r1['number'] in retr and r2['number'] in retr
                and 'no entregado' in retr and 'entregado' in retr and 'UNA vez' in retr, retr[:200])
            await pg.screenshot(path=os.path.join(OUT, 'informe_reenvios.png'), full_page=False)
            for _ in range(3):
                if await pg.locator('#pushLater').count() and await pg.is_visible('#pushLater'): await pg.click('#pushLater'); await pg.wait_for_timeout(300)
            async with pg.expect_download() as dl: await pg.click('#rpExcel')
            f = os.path.join(OUT, 'informe_reenvios.xlsx'); await (await dl.value).save_as(f); wb = load_workbook(f)
            tot = next(r for r in ([c.value for c in x] for x in wb['Resumen'].iter_rows()) if r and r[0] == 'TOTAL')
            rows = [[c.value for c in x] for x in wb['Reintentos'].iter_rows()] if 'Reintentos' in wb.sheetnames else []
            res('RE-E2E-06 Excel = pantalla (envíos e importe) y hoja "Reintentos" con cada intento', tot[1] == n2 and round(tot[6]) == m2
                and any(r and r[0] == 'RE-1001' and r[4] == r2['number'] for r in rows) and any(r and r[0] == 'RE-1001' and r[4] == r1['number'] for r in rows), {'excel': tot, 'reintentos': rows[:3]})
            bc = await pg.evaluate('''async()=>{ const r=await window.OKD_ONLINE.api.reports('preview', { sel: { kind: '7D' } }); const rep=r.report||r; return (rep.best_clients||[]).filter(c=>c.client==='Cliente Ausente').map(c=>c.purchases); }''')
            res('RE-E2E-07 "Mejores clientes": el reintento no suma otra compra', bc in ([1], []), bc)
          finally: await b.close()
    finally:
        if srv.poll() is None: srv.terminate(); srv.wait(5)
    print('TOTAL', sum(1 for r in R if r[1] == 'PASS'), 'PASS /', sum(1 for r in R if r[1] == 'FAIL'), 'FAIL · esperados', EXPECTED)
    return R
if __name__ == '__main__':
    if len(sys.argv) > 1: cr.DIST = sys.argv[1]
    try: rr = asyncio.run(main())
    except Exception as e: traceback.print_exc(); print('ERROR:', repr(e)); sys.exit(2)
    if len(rr) != EXPECTED: print('ERROR: se ejecutaron %d de %d' % (len(rr), EXPECTED)); sys.exit(2)
    sys.exit(1 if any(r[1] != 'PASS' for r in rr) else 0)
